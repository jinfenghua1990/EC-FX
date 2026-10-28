"""商品「大类」top_category：与「分类」category 并存的两层归类。

用户 2026-10-28 明确：现有的分类要保留（后续有用），新增的大类单独一列、自己维护，
报表（对账单导出）里要带上这两列。

这里锁死三条底线：
1. 大类与分类互不覆盖 —— 写大类绝不动分类，反之亦然；
2. 批量写大类只对「值真的变了」的行计数，空转不重复写库；
3. 对账单明细与导出 Excel 里都能拿到大类/分类。
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models import FxDistributor, FxOrder, FxOrderLine, FxProduct
from app.services import distribution as dist


@pytest.fixture()
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    yield session
    session.close()


@pytest.fixture()
def products(db):
    a = FxProduct(sku="A-1", name="百奇巧克力", brand="格力高", category="百奇")
    b = FxProduct(sku="B-1", name="桂格麦片", brand="桂格", category="麦片")
    db.add_all([a, b])
    db.commit()
    return a, b


def test_top_category_and_category_are_independent(db, products):
    a, b = products
    # 批量写大类：分类必须原样保留
    res = dist.set_top_categories(db, [
        {"product_id": a.id, "top_category": "饼干"},
        {"product_id": b.id, "top_category": "冲饮"},
    ])
    assert res["updated"] == 2 and res["errors"] == []
    db.refresh(a); db.refresh(b)
    assert (a.category, a.top_category) == ("百奇", "饼干")
    assert (b.category, b.top_category) == ("麦片", "冲饮")
    view = dist._product_view(a)
    assert view["category"] == "百奇" and view["top_category"] == "饼干"

    # 反过来改分类，大类也不能被抹掉
    dist.update_product(db, a.id, {"category": "百奇分享"})
    db.refresh(a)
    assert (a.category, a.top_category) == ("百奇分享", "饼干")


def test_set_top_category_counts_only_real_changes(db, products):
    a, b = products
    dist.set_top_categories(db, [{"product_id": a.id, "top_category": "饼干"}])
    # 第二次传同样的值 = 没变化，不该算更新
    res = dist.set_top_categories(db, [
        {"product_id": a.id, "top_category": "饼干"},
        {"product_id": b.id, "top_category": "冲饮"},
    ])
    assert res["updated"] == 1
    # 传空字符串 = 清空大类
    dist.set_top_categories(db, [{"product_id": a.id, "top_category": ""}])
    db.refresh(a)
    assert a.top_category == "" and a.category == "百奇"


def test_set_top_category_reports_unknown_product(db):
    res = dist.set_top_categories(db, [{"product_id": 999999, "top_category": "饼干"}])
    assert res["updated"] == 0
    assert res["errors"] and "商品不存在" in res["errors"][0]["message"]


def test_list_top_categories_distinct_sorted(db, products):
    a, b = products
    dist.set_top_categories(db, [
        {"product_id": a.id, "top_category": "冲饮"},
        {"product_id": b.id, "top_category": "饼干"},
    ])
    # 去重 + 排序（SQL 的 order by 按编码序，不是拼音序，断言跟着用 sorted）
    assert dist.list_top_categories(db) == sorted(["饼干", "冲饮"])


@pytest.fixture()
def order_with_lines(db, products):
    a, b = products
    dist.set_top_categories(db, [
        {"product_id": a.id, "top_category": "饼干"},
        {"product_id": b.id, "top_category": "冲饮"},
    ])
    d = FxDistributor(code="FX001", name="测试分销")
    db.add(d); db.commit()
    o = FxOrder(
        order_no="FX20261028-0001", distributor_id=d.id, status="approved",
        order_date=datetime(2026, 10, 20, 10, 0), total_amount=Decimal("100"),
    )
    db.add(o); db.commit()
    db.add_all([
        FxOrderLine(order_id=o.id, product_id=a.id, goods_no="A-1", goods_name="百奇巧克力",
                    quantity=Decimal("2"), unit_price=Decimal("10"), subtotal=Decimal("20")),
        FxOrderLine(order_id=o.id, product_id=b.id, goods_no="B-1", goods_name="桂格麦片",
                    quantity=Decimal("1"), unit_price=Decimal("80"), subtotal=Decimal("80")),
    ])
    db.commit()
    return d, o


def test_statement_lines_carry_top_category(db, order_with_lines):
    d, o = order_with_lines
    data = dist.list_statements(db, distributor_id=d.id, month="2026-10")
    lines = data["orders"][0]["lines"]
    cats = {ln["goods_no"]: (ln["top_category"], ln["category"]) for ln in lines}
    assert cats == {"A-1": ("饼干", "百奇"), "B-1": ("冲饮", "麦片")}


def test_statement_excel_has_top_category_columns(db, order_with_lines):
    d, o = order_with_lines
    filename, blob = dist.statement_excel_bytes(db, distributor_id=d.id, month="2026-10")
    assert filename.endswith("2026-10.xlsx")
    ws = load_workbook(io.BytesIO(blob)).active
    header = [c.value for c in ws[3]]
    assert header[:4] == ["订单号", "下单时间", "大类", "分类"]
    row = [c.value for c in ws[4]]
    assert row[2] in ("饼干", "冲饮")  # 大类
    assert row[3] in ("百奇", "麦片")  # 分类原样带出，不被大类顶掉
