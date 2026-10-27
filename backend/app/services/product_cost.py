"""牌价导入与成本核对状态维护。

牌价（list_price）
------------------
2026-10-27 用户明确：「牌价只有格力高想要，其他品牌不维护」。

来源是品牌订货单 Excel（`20261027抖音红盒子订单.xls` 这类），列结构：
    大类 | 产品代码 | 数量 | 条码 | 类型 | 实开数量 | 缺品数量 | 特价 |
    产品名称 | 规格 | 盒（未税） | 单价(含税) | 订货额(含税) | 实际开单(含税)

换算成「每盒含税牌价」有两条等价路径（实测一致）：
- `单价(含税) ÷ 规格`：209.880324 ÷ 36 = 5.83
- `盒(未税) × 1.13`：5.1593 × 1.13 = 5.83

优先用第一条 —— 表里直接给的两个数做除法，不引入税率假设；
规格缺失时才退回 1.13（格力高全系 13%，但别写死在主路径上）。

牌价只用于算分销折扣率，**不参与利润计算**。利润的成本是采购入库价。
"""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from ..models import FxProduct

# 表头里的列名（容忍空格，表格里常见「产 品 名 称」这种断字）
_COL_BARCODE = "条码"
_COL_SPEC = "规格"
_COL_BOX_EX_TAX = "盒（未税）"
_COL_PRICE_INC_TAX = "单价(含税)"
_COL_NAME = "产品名称"
_VAT_FALLBACK = Decimal("1.13")


def _norm(value: Any) -> str:
    return "".join(str(value or "").split())


def _read_rows(path: str | Path) -> tuple[list[dict], str]:
    """读 Excel，返回 (行列表, sheet 名)。.xls 用 xlrd，.xlsx 用 openpyxl。"""
    file_path = Path(path)
    if not file_path.exists():
        raise ValueError(f"文件不存在：{file_path}")

    suffix = file_path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(file_path))
        sheet = book.sheet_by_index(0)
        grid = [
            [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        sheet_name = sheet.name
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(str(file_path), data_only=True, read_only=True)
        sheet = book.worksheets[0]
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        sheet_name = sheet.title
        book.close()
    else:
        raise ValueError(f"只支持 .xls / .xlsx，收到 {suffix}")

    if not grid:
        raise ValueError("表格是空的")

    header = [_norm(v) for v in grid[0]]
    index: dict[str, int] = {}
    for i, name in enumerate(header):
        for want in (_COL_BARCODE, _COL_SPEC, _COL_BOX_EX_TAX, _COL_PRICE_INC_TAX, _COL_NAME):
            if name == _norm(want) and want not in index:
                index[want] = i

    missing = [c for c in (_COL_BARCODE, _COL_PRICE_INC_TAX) if c not in index]
    if missing:
        raise ValueError(f"表头缺少必要列：{'、'.join(missing)}（实际表头：{' | '.join(header)}）")

    def cell(row: list, col: str) -> Any:
        i = index.get(col)
        return row[i] if i is not None and i < len(row) else ""

    rows: list[dict] = []
    for raw in grid[1:]:
        barcode = str(cell(raw, _COL_BARCODE) or "").strip()
        # Excel 数字单元格会把条码变成 6901845048438.0，去掉尾部 .0
        if barcode.endswith(".0"):
            barcode = barcode[:-2]
        if not barcode or not barcode.isdigit():
            continue
        try:
            spec = Decimal(str(cell(raw, _COL_SPEC) or "0").strip() or "0")
        except Exception:
            spec = Decimal("0")
        try:
            box_ex_tax = Decimal(str(cell(raw, _COL_BOX_EX_TAX) or "0").strip() or "0")
        except Exception:
            box_ex_tax = Decimal("0")
        try:
            price_inc_tax = Decimal(str(cell(raw, _COL_PRICE_INC_TAX) or "0").strip() or "0")
        except Exception:
            price_inc_tax = Decimal("0")

        # 每盒含税牌价：优先 单价(含税)÷规格，其次 盒(未税)×1.13
        if spec > 0 and price_inc_tax > 0:
            unit_price = price_inc_tax / spec
            basis = f"单价(含税){price_inc_tax} ÷ 规格{spec}"
        elif box_ex_tax > 0:
            unit_price = box_ex_tax * _VAT_FALLBACK
            basis = f"盒(未税){box_ex_tax} × {_VAT_FALLBACK}"
        else:
            continue

        rows.append(
            {
                "barcode": barcode,
                "name": str(cell(raw, _COL_NAME) or "").strip(),
                "spec_qty": spec,
                "unit_price": unit_price.quantize(Decimal("0.01")),
                "basis": basis,
                "box_ex_tax": box_ex_tax,
                "price_inc_tax": price_inc_tax,
            }
        )
    return rows, sheet_name


def import_list_prices(
    db: Session, path: str | Path, dry_run: bool = False, overwrite: bool = True
) -> dict[str, Any]:
    """按条码把 Excel 里的牌价写进 fx_products.list_price。

    overwrite=False 时只填「牌价为空」的商品，不动已维护的。
    """
    rows, sheet_name = _read_rows(path)
    barcodes = [r["barcode"] for r in rows]

    products: dict[str, FxProduct] = {}
    for start in range(0, len(barcodes), 400):
        chunk = barcodes[start : start + 400]
        for row in db.scalars(select(FxProduct).where(FxProduct.barcode.in_(chunk))).all():
            if row.barcode:
                products.setdefault(row.barcode, row)

    updated = skipped_existing = 0
    matched: list[dict] = []
    unmatched: list[dict] = []
    for row in rows:
        product = products.get(row["barcode"])
        if product is None:
            unmatched.append({"barcode": row["barcode"], "name": row["name"], "price": str(row["unit_price"])})
            continue
        old = product.list_price or Decimal("0")
        if not overwrite and old > 0:
            skipped_existing += 1
            continue
        if not dry_run:
            product.list_price = row["unit_price"]
        updated += 1
        matched.append(
            {
                "barcode": row["barcode"],
                "sku": product.sku,
                "name": (product.name or "")[:30],
                "list_price": str(row["unit_price"]),
                "old_list_price": str(old),
                "purchase_price": str(product.purchase_price or 0),
                "basis": row["basis"],
            }
        )

    if dry_run:
        db.rollback()
    else:
        db.commit()

    return {
        "dry_run": dry_run,
        "sheet": sheet_name,
        "rows_in_file": len(rows),
        "updated": updated,
        "skipped_existing": skipped_existing,
        "unmatched_count": len(unmatched),
        "unmatched": unmatched[:20],
        "matched_samples": matched[:20],
    }


def set_cost_status(
    db: Session, skus: list[str] | None = None, brand: str = "", status: str = "pending", note: str = ""
) -> dict[str, Any]:
    """标记成本待核对（status="pending"）或恢复正常（status=""）。

    标了 pending 的商品不算毛利 —— 拿错价算利润比不算更糟。
    """
    if status not in {"", "pending"}:
        raise ValueError("status 只支持空字符串或 pending")
    if not skus and not brand.strip():
        raise ValueError("必须指定 skus 或 brand 之一")

    stmt = select(FxProduct)
    clauses = []
    if skus:
        clauses.append(or_(FxProduct.sku.in_(skus), FxProduct.barcode.in_(skus)))
    if brand.strip():
        clauses.append(FxProduct.brand.like(f"%{brand.strip()}%"))
    stmt = stmt.where(or_(*clauses))

    rows = db.scalars(stmt).all()
    for row in rows:
        row.cost_status = status
        row.cost_note = note.strip()[:300] if status else ""
    db.commit()
    return {
        "updated": len(rows),
        "status": status,
        "note": note.strip()[:300] if status else "",
        "skus": [r.sku for r in rows][:50],
    }
