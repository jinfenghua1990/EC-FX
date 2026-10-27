"""表格读取工具：把上传的 XLSX / CSV 读成「表头 → 值」的字典列表。

只保留销售出库导入（`/api/import/wdt-stockout`）需要的读取能力；
原结算单/银行流水导入（依赖 ImportBatch / SettlementRow / BankRow）已随核销模块下线。
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import load_workbook


def _norm(value) -> str:
    return str(value or "").strip().replace("\n", "")


def _rows_from_xlsx(content: bytes) -> list[dict[str, object]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    headers = [_norm(v) for v in next(iterator)]
    return [{headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
            for values in iterator if any(v not in (None, "") for v in values)]


def _rows_from_csv(content: bytes) -> list[dict[str, object]]:
    text = None
    for encoding in ("utf-8-sig", "gb18030", "gbk"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别")
    return list(csv.DictReader(io.StringIO(text)))


def read_tabular(filename: str, content: bytes) -> list[dict[str, object]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return _rows_from_xlsx(content)
    if suffix == ".csv":
        return _rows_from_csv(content)
    raise ValueError("仅支持 XLSX / CSV")
